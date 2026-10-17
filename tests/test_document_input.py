"""Uploaded text must be bounded and preserve DOCX reading order."""
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import app
from ingestion_pipeline.document_input import MAX_FILE_BYTES, extract_text


def test_docx_paragraphs_and_tables_keep_document_order():
    document = Document()
    document.add_paragraph("First term")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Risk"
    table.cell(0, 1).text = "High"
    document.add_paragraph("Last term")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_text("contract.docx", buffer.getvalue())

    assert text.index("First term") < text.index("Risk\tHigh") < text.index("Last term")


def test_upload_rejects_unsupported_or_oversize_content():
    with pytest.raises(ValueError, match="TXT, PDF, or DOCX"):
        extract_text("contract.exe", b"text")
    with pytest.raises(ValueError, match="5 MB"):
        extract_text("contract.txt", b"x" * (MAX_FILE_BYTES + 1))


def test_upload_endpoint_returns_text_without_storing_file():
    response = TestClient(app).post(
        "/api/extract", files={"file": ("contract.txt", b"First clause.\n\nSecond clause.", "text/plain")})

    assert response.status_code == 200
    assert response.json() == {"text": "First clause.\n\nSecond clause."}
