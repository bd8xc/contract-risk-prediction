import re
from pathlib import Path

from pypdf import PdfReader
from docx import Document


class DocumentProcessor:
    """
    Extracts and cleans text from legal documents.

    Supported formats:
        - PDF
        - DOCX
        - TXT

    Responsibilities:
        1. Validate the input file
        2. Extract text
        3. Preserve useful legal-document structure
        4. Clean common PDF extraction artifacts
        5. Return text + document metadata

    This class does NOT:
        - split clauses
        - classify risk
        - run LegalBERT
    """

    SUPPORTED_EXTENSIONS = {
        ".pdf",
        ".docx",
        ".txt",
    }

    ##########################################################
    # MAIN PROCESSING METHOD
    ##########################################################

    def process(self, file_path):
        """
        Process a PDF, DOCX, or TXT file.

        Returns:

        {
            "file_name": str,
            "file_type": str,
            "text": str,
            "pages": list,
            "page_count": int | None,
            "character_count": int,
            "word_count": int
        }
        """

        file_path = Path(file_path)

        ######################################################
        # CHECK FILE EXISTS
        ######################################################

        if not file_path.exists():

            raise FileNotFoundError(
                f"File not found: {file_path}"
            )

        ######################################################
        # CHECK FILE TYPE
        ######################################################

        extension = (
            file_path
            .suffix
            .lower()
        )

        if extension not in self.SUPPORTED_EXTENSIONS:

            supported = ", ".join(
                sorted(
                    self.SUPPORTED_EXTENSIONS
                )
            )

            raise ValueError(
                f"Unsupported file type: {extension}\n"
                f"Supported types: {supported}"
            )

        ######################################################
        # EXTRACT BASED ON TYPE
        ######################################################

        if extension == ".pdf":

            result = self._extract_pdf(
                file_path
            )

        elif extension == ".docx":

            result = self._extract_docx(
                file_path
            )

        elif extension == ".txt":

            result = self._extract_txt(
                file_path
            )

        else:

            raise ValueError(
                f"Unsupported file type: {extension}"
            )

        ######################################################
        # CLEAN COMPLETE DOCUMENT
        ######################################################

        cleaned_text = self.clean_text(
            result["text"]
        )

        result["text"] = cleaned_text

        ######################################################
        # ADD METADATA
        ######################################################

        result["file_name"] = (
            file_path.name
        )

        result["file_type"] = (
            extension
        )

        result["character_count"] = len(
            cleaned_text
        )

        result["word_count"] = len(
            cleaned_text.split()
        )

        return result

    ##########################################################
    # PDF EXTRACTION
    ##########################################################

    def _extract_pdf(
        self,
        file_path,
    ):

        reader = PdfReader(
            str(file_path)
        )

        pages = []

        all_text = []

        ######################################################
        # PROCESS EACH PAGE
        ######################################################

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):

            try:

                text = (
                    page.extract_text(
                    extraction_mode="layout"
                            )
                            or ""
                        )

            except Exception as error:

                print(
                    f"Warning: Could not extract "
                    f"page {page_number}: {error}"
                )

                text = ""

            ##################################################
            # CLEAN PAGE TEXT
            ##################################################

            cleaned_page = self.clean_text(
                text
            )

            ##################################################
            # SAVE PAGE INFORMATION
            ##################################################

            pages.append(
                {
                    "page_number":
                        page_number,

                    "text":
                        cleaned_page,
                }
            )

            if cleaned_page:

                all_text.append(
                    cleaned_page
                )

        ######################################################
        # COMBINE PAGES
        ######################################################

        combined_text = "\n\n".join(
            all_text
        )

        ######################################################
        # POSSIBLE SCANNED PDF DETECTION
        ######################################################

        if len(
            combined_text.strip()
        ) < 50:

            print(
                "\nWARNING:"
            )

            print(
                "Very little text was extracted "
                "from this PDF."
            )

            print(
                "The document may be scanned "
                "and may require OCR."
            )

        return {

            "text":
                combined_text,

            "pages":
                pages,

            "page_count":
                len(
                    reader.pages
                ),
        }

    ##########################################################
    # DOCX EXTRACTION
    ##########################################################

    def _extract_docx(
        self,
        file_path,
    ):

        document = Document(
            str(file_path)
        )

        blocks = []

        ######################################################
        # EXTRACT PARAGRAPHS
        ######################################################

        for paragraph in document.paragraphs:

            text = (
                paragraph
                .text
                .strip()
            )

            if text:

                blocks.append(
                    text
                )

        ######################################################
        # EXTRACT TABLE CONTENT
        ######################################################

        for table in document.tables:

            for row in table.rows:

                row_values = []

                for cell in row.cells:

                    cell_text = (
                        cell
                        .text
                        .strip()
                    )

                    if cell_text:

                        row_values.append(
                            cell_text
                        )

                if row_values:

                    blocks.append(
                        " | ".join(
                            row_values
                        )
                    )

        ######################################################
        # COMBINE CONTENT
        ######################################################

        text = "\n\n".join(
            blocks
        )

        return {

            "text":
                text,

            # DOCX does not provide reliable
            # page boundaries through python-docx.

            "pages":
                [],

            "page_count":
                None,
        }

    ##########################################################
    # TXT EXTRACTION
    ##########################################################

    def _extract_txt(
        self,
        file_path,
    ):

        ######################################################
        # TRY UTF-8 FIRST
        ######################################################

        try:

            with open(
                file_path,
                "r",
                encoding="utf-8",
            ) as file:

                text = file.read()

        ######################################################
        # FALLBACK FOR OLDER TEXT FILES
        ######################################################

        except UnicodeDecodeError:

            with open(
                file_path,
                "r",
                encoding="latin-1",
            ) as file:

                text = file.read()

        return {

            "text":
                text,

            "pages":
                [],

            "page_count":
                None,
        }

    ##########################################################
    # TEXT CLEANING
    ##########################################################

    @staticmethod
    def clean_text(text):
        """
        Conservative cleaning for legal documents.

        Important:
        We preserve line structure because numbered sections
        such as:

            7. TERMINATION
            7.1 Either party...
            7.2 The Supplier...

        are needed later by ClauseSegmenter.
        """

        if not text:

            return ""

        ######################################################
        # NORMALIZE LINE ENDINGS
        ######################################################

        text = text.replace(
            "\r\n",
            "\n",
        )

        text = text.replace(
            "\r",
            "\n",
        )

        ######################################################
        # REMOVE NULL CHARACTERS
        ######################################################

        text = text.replace(
            "\x00",
            "",
        )

        ######################################################
        # NORMALIZE NON-BREAKING SPACES
        ######################################################

        text = text.replace(
            "\u00a0",
            " ",
        )

        ######################################################
        # NORMALIZE OTHER UNUSUAL SPACES
        ######################################################

        text = text.replace(
            "\u2007",
            " ",
        )

        text = text.replace(
            "\u202f",
            " ",
        )

        ######################################################
        # FIX HYPHENATED WORDS ACROSS LINES
        #
        # Example:
        #
        # indemni-
        # fication
        #
        # becomes:
        #
        # indemnification
        ######################################################

        text = re.sub(
            r"([A-Za-z])-\s*\n\s*([A-Za-z])",
            r"\1\2",
            text,
        )

        ######################################################
        # CLEAN EACH LINE INDIVIDUALLY
        #
        # Important:
        # We do NOT remove newlines here.
        ######################################################

        lines = text.split(
            "\n"
        )

        cleaned_lines = []

        for line in lines:

            ##################################################
            # COLLAPSE MULTIPLE SPACES/TABS
            ##################################################

            line = re.sub(
                r"[ \t]+",
                " ",
                line,
            )

            ##################################################
            # REMOVE LEADING/TRAILING WHITESPACE
            ##################################################

            line = line.strip()

            cleaned_lines.append(
                line
            )

        ######################################################
        # REBUILD DOCUMENT
        ######################################################

        text = "\n".join(
            cleaned_lines
        )

        ######################################################
        # LIMIT EXCESSIVE BLANK LINES
        ######################################################

        text = re.sub(
            r"\n{3,}",
            "\n\n",
            text,
        )

        return text.strip()


##############################################################
# COMMAND-LINE TEST
##############################################################

if __name__ == "__main__":

    processor = DocumentProcessor()

    print(
        "=" * 70
    )

    print(
        "LEGAL DOCUMENT PROCESSOR"
    )

    print(
        "=" * 70
    )

    print(
        "\nEnter the path to a "
        "PDF, DOCX, or TXT contract."
    )

    print(
        "Type 'quit' to exit.\n"
    )

    while True:

        ######################################################
        # GET FILE PATH
        ######################################################

        file_path = input(
            "Document path: "
        ).strip()

        ######################################################
        # REMOVE WINDOWS QUOTES
        ######################################################

        file_path = file_path.strip(
            "\"'"
        )

        ######################################################
        # EXIT
        ######################################################

        if file_path.lower() in {
            "quit",
            "exit",
            "q",
        }:

            print(
                "\nExiting."
            )

            break

        ######################################################
        # PROCESS DOCUMENT
        ######################################################

        try:

            result = processor.process(
                file_path
            )

            print(
                "\n"
                + "=" * 70
            )

            print(
                "DOCUMENT EXTRACTED SUCCESSFULLY"
            )

            print(
                "=" * 70
            )

            ##################################################
            # METADATA
            ##################################################

            print(
                f"\nFile name       : "
                f"{result['file_name']}"
            )

            print(
                f"File type       : "
                f"{result['file_type']}"
            )

            print(
                f"Pages           : "
                f"{result['page_count']}"
            )

            print(
                f"Characters      : "
                f"{result['character_count']}"
            )

            print(
                f"Words           : "
                f"{result['word_count']}"
            )

            ##################################################
            # PREVIEW
            ##################################################

            print(
                "\n--- TEXT PREVIEW ---\n"
            )

            preview_length = 3000

            preview = (
                result["text"][
                    :preview_length
                ]
            )

            print(
                preview
            )

            if (
                len(
                    result["text"]
                )
                > preview_length
            ):

                print(
                    "\n\n"
                    "[Preview truncated]"
                )

            print(
                "\n"
                + "=" * 70
            )

        ######################################################
        # ERROR
        ######################################################

        except Exception as error:

            print(
                f"\nProcessing failed: "
                f"{error}\n"
            )